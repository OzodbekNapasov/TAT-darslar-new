# -*- coding: utf-8 -*-
"""
Word 2-dars amaliy ishi: 6 variantli tarqatma PDF, etalon DOCX va dars sahifasidagi variant kartalari.

Ishga tushirish (shu papkada):  python build.py
Natija:
  ../tarqatma/N-variant.pdf           — talabaga: 1-bet topshiriq va baholash jadvali, 2-bet namuna (A4)
  ../tarqatma/barcha-variantlar.pdf   — 6 ta variant bitta faylda (chop etish uchun)
  ../tarqatma/baholash-qaydnomasi.pdf — o'qituvchiga: tekshirish tartibi va guruh qaydnomasi
  ../tarqatma/etalon/N-variant-etalon.docx — o'qituvchiga: to'g'ri formatlangan Word hujjat
  ../dars.html ichidagi VARIANTS bloki
Matnni o'zgartirish uchun variants.py ni tahrirlang va qayta ishga tushiring.
"""
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from variants import CRITERIA, GRADES, TEXNIKUM, VARIANTS, interval_key, interval_str

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT_HTML = os.path.join(HERE, 'out')
PDF_DIR = os.path.join(ROOT, 'tarqatma')
DOCX_DIR = os.path.join(PDF_DIR, 'etalon')
LESSON = os.path.join(ROOT, 'dars.html')
CHROME = r'C:\Program Files\Google\Chrome\Application\chrome.exe'

with open(os.path.join(HERE, 'sheet.css'), encoding='utf8') as f:
    SHEET_CSS = f.read()


def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def kc(s):
    """[[Ctrl]] -> tugma ko'rinishi."""
    return re.sub(r'\[\[(.+?)\]\]', r'<span class="key-cap">\1</span>', s)


# ----------------------------------------------------------------------------
# NAMUNA VARAG'I (haqiqiy A4: maydonlar 3 / 1,5 / 2 / 2 sm, Times New Roman 14 pt)
# ----------------------------------------------------------------------------
def sample_body(V):
    out = []
    for b in V['blocks']:
        k = b[0]
        if k == 'gap':
            out.append('<div class="gap"></div>')
        elif k in ('right', 'sign'):
            out += ['<p class="r">%s</p>' % esc(x) for x in b[1]]
        elif k == 'title':
            out.append('<p class="c b">%s</p>' % esc(b[1]))
        elif k == 'sub':
            out.append('<p class="c">%s</p>' % esc(b[1]))
        elif k == 'p':
            out.append('<p class="j ind">%s</p>' % esc(b[1]))
        elif k == 'left':
            out.append('<p>%s</p>' % esc(b[1]))
    return '\n'.join(out)


def sample_sheet(V, label=True):
    lab = '<span class="sheet-label">Namuna • %d-variant</span>' % V['n'] if label else ''
    cls = 'ls15' if V['interval'] == 1.5 else 'ls10'
    return '<div class="sheet doc %s">%s%s</div>' % (cls, lab, sample_body(V))


# ----------------------------------------------------------------------------
# TOPSHIRIQ VARAG'I (talaba uchun 1-bet)
# ----------------------------------------------------------------------------
def steps(V):
    kinds = [b[0] for b in V['blocks']]
    n_right = sum(len(b[1]) for b in V['blocks'] if b[0] == 'right')
    n_sign = sum(len(b[1]) for b in V['blocks'] if b[0] == 'sign')
    s = [
        'Word’ni oching → <b>Yangi hujjat</b>. Darhol saqlang: [[Ctrl]]+[[S]] → o‘qituvchi ko‘rsatgan papka → '
        'nom: <b>Familiya_V%d</b> (masalan, <i>Karimov_V%d</i>).' % (V['n'], V['n']),
        '<b>Maket → Maydonlar → Maxsus maydonlar</b>: yuqori 2 sm, pastki 2 sm, chap 3 sm, o‘ng 1,5 sm.',
        '<b>Bosh sahifa → Shrift</b>: Times New Roman, 14 pt.',
        'Namunadagi matnni xuddi o‘zidek tering. [[Enter]] — faqat abzats oxirida va namunadagi bo‘sh qator o‘rnida. '
        'So‘zlar orasida bitta probel, tinish belgisidan <u>keyin</u> probel. o‘, g‘ dagi belgi: [[Alt]]+[[0145]].',
    ]
    if 'right' in kinds:
        s.append('«Kimga» qismini (yuqoridagi %d qator) sichqonchani <b>bosib turib tortish</b> bilan belgilang → '
                 '[[Ctrl]]+[[R]] (o‘ngga).' % n_right)
    t = 'Sarlavhani belgilang: qatorning <b>chap chetida bir marta bosing</b> → [[Ctrl]]+[[B]], [[Ctrl]]+[[E]].'
    if 'sub' in kinds:
        t += ' Ostidagi izoh qatori ham markazga ([[Ctrl]]+[[E]]), lekin qalin emas.'
    s.append(t)
    s.append('Asosiy matnni belgilang: birinchi abzats boshida bosing, [[Shift]]ni bosib turib oxirgi abzats '
             'oxirida bosing → [[Ctrl]]+[[J]]. <b>Abzats</b> oynasi → Birinchi qator: <b>Chekinish 1,25 sm</b>.')
    s.append('Imzo qatori%s: kursorni qator boshiga qo‘yib, [[Shift]]+[[End]] bilan belgilang → [[Ctrl]]+[[R]]. '
             '%s' % ('lari (%d ta)' % n_sign if n_sign > 1 else '',
                     'Sana qatori chapda qoladi.' if 'left' in kinds else ''))
    s.append('[[Ctrl]]+[[A]] bilan butun hujjatni belgilang → qatorlar oralig‘i <b>%s</b> ([[Ctrl]]+[[%s]]).'
             % (interval_str(V['interval']), interval_key(V['interval'])))
    s.append('<b>¶</b> tugmasini yoqing ([[Ctrl]]+[[Shift]]+[[8]]): ortiqcha probel (··) va bo‘sh ¶ larni o‘chiring. '
             '[[Ctrl]]+[[S]] bilan saqlang va o‘qituvchiga ayting.')
    return s


def info_sheet(V):
    reqs = [
        ('Varaq', 'A4, kitob'),
        ('Maydonlar', 'yuqori 2, pastki 2, chap 3, o‘ng 1,5 sm'),
        ('Shrift', 'Times New Roman, 14 pt'),
        ('Sarlavha', 'qalin, markazda'),
        ('Asosiy matn', 'ikki tomonga, xatboshi 1,25 sm'),
        ('Qatorlar oralig‘i', interval_str(V['interval'])),
        ('Fayl nomi', 'Familiya_V%d.docx' % V['n']),
    ]
    sel = [
        ('2 marta bosish', 'bitta so‘z'),
        ('3 marta bosish', 'butun abzats'),
        ('Chap chetda bosish', 'butun qator'),
        ('Bosib turib tortish', 'istalgan qism'),
        ('[[Shift]] + bosish', 'kursordan bosilgan joygacha'),
        ('[[Shift]]+[[End]]', 'qator oxirigacha'),
        ('[[Ctrl]]+[[A]]', 'butun hujjat'),
    ]
    crit = ''.join('<tr><td class="cn">%d</td><td><b>%s</b> — %s</td><td class="cb">0,5</td><td class="cg"></td></tr>'
                   % (i + 1, a, b.replace('Familiya_V№', 'Familiya_V%d' % V['n'])) for i, (a, b) in enumerate(CRITERIA))
    grades = '; '.join('%s → %s' % (a, b.split()[0]) for a, b in GRADES)
    return '''<div class="sheet info">
  <div class="info-top"><div><b>Tibbiyotda axborot texnologiyalari</b><br>Word 2-dars: matn kiritish va rasmiy hujjat • amaliy ish</div><span class="info-num">%(n)d-variant</span></div>
  <div class="info-fio"><span>F.I.O.:</span><span>Guruh:</span><span>Sana:</span></div>
  <h1>%(n)d-variant: %(title)s</h1>
  <p class="info-goal"><b>Hujjat:</b> %(about)s Keyingi betdagi namunani Word’da xuddi shunday tering va rasmiy hujjat talablari bo‘yicha formatlang.</p>
  <h2>Topshiriq sharti</h2>
  <ol class="steps">%(steps)s</ol>
  <div class="info-grid">
    <div><h2>Hujjat talablari</h2><table class="kt">%(reqs)s</table></div>
    <div><h2>Belgilash usullari</h2><table class="kt">%(sel)s</table></div>
  </div>
  <h2>Baholash (o‘qituvchi to‘ldiradi)</h2>
  <table class="crit"><thead><tr><th>№</th><th>Mezon</th><th>Maks.</th><th>Ball</th></tr></thead><tbody>%(crit)s
    <tr class="tot"><td></td><td>Jami</td><td class="cb">5</td><td class="cg"></td></tr></tbody></table>
  <div class="info-sign"><span><b>Baho:</b> ______ <span class="gr">(%(grades)s)</span></span><span>O‘qituvchi imzosi: __________</span></div>
  <div class="info-foot"><span>%(tex)s</span><span>1-bet: topshiriq • 2-bet: namuna</span></div>
</div>''' % dict(n=V['n'], title=V['title'], about=V['about'],
                 steps=''.join('<li>%s</li>' % kc(x) for x in steps(V)),
                 reqs=''.join('<tr><td>%s</td><td><b>%s</b></td></tr>' % r for r in reqs),
                 sel=''.join('<tr><td>%s</td><td>%s</td></tr>' % (kc(a), b) for a, b in sel),
                 crit=crit, grades=grades, tex=TEXNIKUM)


# ----------------------------------------------------------------------------
# O'QITUVCHI: TEKSHIRISH TARTIBI VA QAYDNOMA
# ----------------------------------------------------------------------------
def teacher_sheets():
    check = [
        ('10', 'Fayl nomini ko‘ring: <b>Familiya_V№.docx</b>. Variant raqami talabaga berilgan tarqatma bilan bir xil bo‘lishi kerak.'),
        ('1', '<b>Maket → Maydonlar → Maxsus maydonlar</b>: 2 / 2 / 3 / 1,5 sm.'),
        ('2', '[[Ctrl]]+[[A]] bosing: Shrift maydonida <b>Times New Roman</b> va <b>14</b> ko‘rinsa — to‘g‘ri. Maydon <b>bo‘sh</b> qolsa, hujjatda boshqa shrift yoki o‘lcham aralashgan.'),
        ('3, 4', '<b>Taqriz → Solishtirish → Solishtirish</b> (Review → Compare): asl hujjat — etalon <code>N-variant-etalon.docx</code>, o‘zgartirilgan — talaba fayli. Tushib qolgan, ortiqcha va xato yozilgan so‘zlar rangli bo‘lib chiqadi.'),
        ('5, 6, 7', '<b>¶</b> tugmasini yoqing ([[Ctrl]]+[[Shift]]+[[8]]): ortiqcha ¶ (qator o‘rtasida Enter, ketma-ket bo‘sh qatorlar), ikki va undan ko‘p nuqta (··), tinish belgisi oldidagi nuqta, qavs ichidagi nuqta.'),
        ('8', 'Sarlavha qalin va markazda; «kimga» va imzo o‘ngda, sana chapda (namunadagidek).'),
        ('9, 10', 'Asosiy abzats ichida bosing → <b>Abzats</b> oynasi: tekislash «Ikki tomonga», birinchi qator «Chekinish 1,25 sm», interval variantdagidek (toq variantlar 1,5; juft variantlar 1,0).'),
    ]
    rules = [
        'Har bir mezon to‘liq bajarilsa 0,5 ball, bajarilmasa 0 ball. Yarim ball (0,25) qo‘yilmaydi — tekshirish tez va bir xil bo‘ladi.',
        '3-mezon: 1–2 ta so‘z tushib qolsa ham 0 ball. 4-mezon: 3 tadan ortiq imlo xatosi bo‘lsa 0 ball.',
        'Probel bilan surilgan qator (masalan, «kimga» qismi probellar bilan o‘ngga surilgan) — 6- va 8-mezonlar 0 ball.',
        'Tire (—) o‘rniga defis (-) yoki chiziqcha (–) yozilsa, qo‘shtirnoq “ ” o‘rniga " " yozilsa — xato hisoblanmaydi.',
        'o‘, g‘ dagi belgi ‘ , ’ yoki \' bo‘lishi mumkin — xato emas. Belgisiz (“ogil”) yoki ` bilan (“o`g`il”) yozilsa — imlo xatosi.',
        '«Kimga» qatorlari [[Enter]] yoki [[Shift]]+[[Enter]] bilan ajratilishi mumkin — ikkalasi ham to‘g‘ri.',
        'Ikki talabaning fayli bir xil bo‘lsa (variant raqami boshqa bo‘lsa ham) — ikkalasiga ham qayta topshiriq beriladi.',
    ]
    page1 = '''<div class="sheet info teacher">
  <div class="info-top"><div><b>Tibbiyotda axborot texnologiyalari</b><br>Word 2-dars amaliy ishi • faqat o‘qituvchi uchun</div><span class="info-num">O‘qituvchi</span></div>
  <h1 style="margin-top:0.4cm">Amaliy ishni tekshirish tartibi</h1>
  <p class="info-goal">6 ta variant: toq variantlarda (1, 3, 5) qatorlar oralig‘i <b>1,5</b>, juft variantlarda (2, 4, 6) — <b>1,0</b>. Yonma-yon o‘tirgan talabalarga turli variant bering: ko‘chirilgan fayl darhol bilinadi.</p>
  <h2>Qadamlar (bitta ishga 2–3 daqiqa)</h2>
  <table class="crit chk"><thead><tr><th>Mezon</th><th>Qanday tekshiriladi</th></tr></thead><tbody>%s</tbody></table>
  <h2>Mezonlar</h2>
  <table class="crit"><thead><tr><th>№</th><th>Mezon</th><th>Ball</th></tr></thead><tbody>%s</tbody></table>
  <p class="info-goal" style="margin-top:0.15cm">%s</p>
  <h2>Baholash qoidalari</h2>
  <ul class="rules">%s</ul>
  <div class="info-foot"><span>%s</span><span>1-bet: tekshirish tartibi • 2-bet: guruh qaydnomasi</span></div>
</div>''' % (
        ''.join('<tr><td class="cn">%s</td><td>%s</td></tr>' % (a, kc(b)) for a, b in check),
        ''.join('<tr><td class="cn">%d</td><td><b>%s</b> — %s</td><td class="cb">0,5</td></tr>' % (i + 1, a, b) for i, (a, b) in enumerate(CRITERIA)),
        'Jami: 5 ball. ' + ' &nbsp;•&nbsp; '.join('%s → <b>%s</b>' % g for g in GRADES),
        ''.join('<li>%s</li>' % kc(r) for r in rules),
        TEXNIKUM)

    head = ''.join('<th>%d</th>' % i for i in range(1, 11))
    rows = ''.join('<tr><td class="cn">%d</td><td></td><td></td>%s<td></td><td></td></tr>' % (i, '<td></td>' * 10) for i in range(1, 31))
    page2 = '''<div class="sheet info teacher">
  <div class="info-top"><div><b>Tibbiyotda axborot texnologiyalari</b><br>Word 2-dars amaliy ishi • baholash qaydnomasi</div><span class="info-num">Qaydnoma</span></div>
  <div class="info-fio"><span>Guruh:</span><span>Sana:</span><span>O‘qituvchi:</span></div>
  <table class="reg"><thead><tr><th>№</th><th class="fio">F.I.O.</th><th class="w1">Var.</th>%s<th class="w2">Jami</th><th class="w2">Baho</th></tr></thead><tbody>%s</tbody></table>
  <p class="reg-legend">1 — maydonlar; 2 — shrift; 3 — matn to‘liq; 4 — imlo; 5 — Enter; 6 — probel; 7 — tinish belgilari; 8 — sarlavha va tekislash; 9 — asosiy matn; 10 — interval va fayl. Har bir mezon 0,5 ball.</p>
  <div class="info-foot"><span>%s</span><span>Imzo: ______________</span></div>
</div>''' % (head, rows, TEXNIKUM)
    return [page1, page2]


EXTRA_CSS = '''
  .sheet-label { position: absolute; top: 0.35cm; right: 1.5cm; font-family: Arial, Helvetica, sans-serif; font-size: 8pt; color: #a6a6a6; letter-spacing: 0.05em; text-transform: uppercase; }
  .sheet.doc.ls10 { line-height: 1.15; }
  .sheet.doc.ls10 .gap { height: 1.15em; }
  .sheet.info { font-size: 9.6pt; line-height: 1.36; }
  .sheet.info h1 { font-size: 15pt; }
  .sheet.info h2 { margin: 0.28cm 0 0.12cm 0; font-size: 10pt; }
  .sheet.info ol.steps li { margin-bottom: 0.04cm; }
  .sheet.info ol.steps b, .sheet.info table b { color: #0b2a5b; }
  .sheet.info table.kt td { padding: 2pt 4pt; font-size: 9pt; }
  .sheet.info table.kt td:first-child { width: 40%; }
  table.crit { width: 100%; border-collapse: collapse; font-size: 8.8pt; line-height: 1.25; }
  table.crit th { background: #185abd; color: #fff; font-weight: 700; padding: 2.5pt 4pt; text-align: left; border: 0.75pt solid #185abd; }
  table.crit td { border: 0.75pt solid #c9d6ea; padding: 1.4pt 4pt; vertical-align: top; }
  table.crit td.cn { width: 0.7cm; text-align: center; font-weight: 700; color: #185abd; }
  table.crit td.cb { width: 1.2cm; text-align: center; }
  table.crit td.cg { width: 1.4cm; }
  table.crit tr.tot td { font-weight: 700; background: #eef3fb; }
  table.crit .gr { font-weight: 400; font-size: 8pt; color: #333; }
  table.crit.chk td.cn { width: 1.4cm; }
  .info-sign { display: flex; justify-content: space-between; align-items: baseline; margin-top: 0.25cm; font-size: 10pt; }
  .info-sign .gr { font-size: 8pt; color: #555; }
  .sheet.info .info-fio { margin: 0.25cm 0 0.3cm 0; }
  .sheet.info ul.rules { padding-left: 0.5cm; }
  .sheet.info ul.rules li { margin-bottom: 0.08cm; }
  .sheet.info code { font-family: Consolas, "Courier New", monospace; font-size: 8.5pt; }
  table.reg { width: 100%; border-collapse: collapse; table-layout: fixed; font-size: 8.5pt; margin-top: 0.2cm; }
  table.reg th { width: 0.72cm; }
  table.reg th.w1 { width: 1.1cm; } table.reg th.w2 { width: 1.4cm; }
  table.reg th { background: #185abd; color: #fff; border: 0.75pt solid #185abd; padding: 3pt 0; text-align: center; font-weight: 700; }
  table.reg td { border: 0.75pt solid #9fb4d6; height: 0.72cm; text-align: center; }
  table.reg th.fio { width: auto; }
  table.reg td.cn { width: 0.7cm; color: #185abd; font-weight: 700; }
  .reg-legend { font-size: 8.5pt; color: #333; margin-top: 0.2cm !important; }
'''


def print_doc(sheets, title):
    return '''<!DOCTYPE html>
<html lang="uz"><head><meta charset="UTF-8"><title>%s</title>
<style>
  @page { size: A4; margin: 0; }
  html, body { margin: 0; padding: 0; background: #ffffff; }
  .sheet { page-break-after: always; break-after: page; }
  .sheet:last-child { page-break-after: auto; break-after: auto; }
%s%s</style></head><body>
%s
</body></html>''' % (title, SHEET_CSS, EXTRA_CSS, '\n'.join(sheets))


def to_pdf(html_path, pdf_path):
    subprocess.run([CHROME, '--headless=new', '--disable-gpu', '--no-pdf-header-footer',
                    '--print-to-pdf=' + pdf_path, 'file:///' + html_path.replace('\\', '/')],
                   check=True, capture_output=True)
    with open(pdf_path, 'rb') as f:
        return len(re.findall(rb'/Type\s*/Page[^s]', f.read()))


# ----------------------------------------------------------------------------
# ETALON DOCX (o'qituvchi uchun: Taqriz -> Solishtirish bilan tekshirishga)
# ----------------------------------------------------------------------------
def etalon_docx(V, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)

    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(14)
    rpr = st.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(a), 'Times New Roman')
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE if V['interval'] == 1.5 else WD_LINE_SPACING.SINGLE

    def para(text='', align=None, bold=False, indent=False):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        if indent:
            p.paragraph_format.first_line_indent = Cm(1.25)
        if text:
            r = p.add_run(text)
            r.bold = bold or None
        return p

    for b in V['blocks']:
        k = b[0]
        if k == 'gap':
            para()
        elif k in ('right', 'sign'):
            for x in b[1]:
                para(x, WD_ALIGN_PARAGRAPH.RIGHT)
        elif k == 'title':
            para(b[1], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        elif k == 'sub':
            para(b[1], WD_ALIGN_PARAGRAPH.CENTER)
        elif k == 'p':
            para(b[1], WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True)
        elif k == 'left':
            para(b[1], WD_ALIGN_PARAGRAPH.LEFT)

    cp = doc.core_properties
    cp.title = '%d-variant: %s (etalon)' % (V['n'], V['name'])
    cp.author = TEXNIKUM
    cp.comments = 'Word 2-dars amaliy ishi uchun to‘g‘ri formatlangan namuna.'
    doc.save(path)


# ----------------------------------------------------------------------------
# DARS SAHIFASIDAGI VARIANT KARTALARI
# ----------------------------------------------------------------------------
START, END = '<!-- VARIANTS_START -->', '<!-- VARIANTS_END -->'


def variant_cards():
    cards = []
    for V in VARIANTS:
        cards.append('''        <article class="var-card" style="--vc:%(c)s">
          <div class="var-top"><span class="var-num">%(n)d-variant</span><span class="var-int">Interval %(iv)s</span></div>
          <h3>%(name)s</h3>
          <p>%(about)s</p>
          <button type="button" class="var-preview-btn" data-variant="%(n)d">Namunani ko‘rish</button>
          <div class="var-actions">
            <a href="./tarqatma/%(n)d-variant.pdf" class="nav-btn" download>PDF yuklab olish</a>
            <a href="./tarqatma/etalon/%(n)d-variant-etalon.docx" class="nav-btn teacher-only" download>Etalon DOCX</a>
          </div>
          <template class="var-sample">%(sample)s</template>
        </article>''' % dict(n=V['n'], name=V['name'], about=V['about'], iv=interval_str(V['interval']),
                             c=['#60a5fa', '#34d399', '#f59e0b', '#a78bfa', '#f472b6', '#22d3ee'][V['n'] - 1],
                             sample=sample_sheet(V, label=False)))
    return '%s\n        <div class="var-grid">\n%s\n        </div>\n        %s' % (START, '\n'.join(cards), END)


def main():
    os.makedirs(OUT_HTML, exist_ok=True)
    os.makedirs(DOCX_DIR, exist_ok=True)

    def write(path, text):
        with open(path, 'w', encoding='utf8', newline='\n') as f:
            f.write(text)

    ok = True
    all_sheets = []
    for V in VARIANTS:
        sheets = [info_sheet(V), sample_sheet(V)]
        all_sheets += sheets
        src = os.path.join(OUT_HTML, '%d-variant.html' % V['n'])
        write(src, print_doc(sheets, '%d-variant: %s' % (V['n'], V['name'])))
        pages = to_pdf(src, os.path.join(PDF_DIR, '%d-variant.pdf' % V['n']))
        etalon_docx(V, os.path.join(DOCX_DIR, '%d-variant-etalon.docx' % V['n']))
        ok &= pages == 2
        print('%d-variant (%s): %d bet %s' % (V['n'], V['name'], pages, 'OK' if pages == 2 else 'DIQQAT: 2 bet kutilgan'))

    src = os.path.join(OUT_HTML, 'barcha-variantlar.html')
    write(src, print_doc(all_sheets, 'Word 2-dars: barcha variantlar'))
    pages = to_pdf(src, os.path.join(PDF_DIR, 'barcha-variantlar.pdf'))
    ok &= pages == len(all_sheets)
    print('Barcha variantlar: %d bet (kutilgan %d)' % (pages, len(all_sheets)))

    src = os.path.join(OUT_HTML, 'baholash-qaydnomasi.html')
    write(src, print_doc(teacher_sheets(), 'Word 2-dars: baholash qaydnomasi'))
    pages = to_pdf(src, os.path.join(PDF_DIR, 'baholash-qaydnomasi.pdf'))
    ok &= pages == 2
    print('Baholash qaydnomasi: %d bet' % pages)

    if os.path.exists(LESSON):
        with open(LESSON, encoding='utf8', newline='') as f:
            html = f.read()
        if START in html:
            html = re.sub(re.escape(START) + r'.*?' + re.escape(END), lambda m: variant_cards(), html, flags=re.S)
            write(LESSON, html)
            print('dars.html: variant kartalari yangilandi')

    if not ok:
        sys.exit(1)


if __name__ == '__main__':
    main()
